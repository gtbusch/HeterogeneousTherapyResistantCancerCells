import pandas as pd
import numpy as np
import openpyxl
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import ColorScaleRule

# =============================================================================
# HELPER FUNCTIONS
# =============================================================================

def normalise_wgs(v):
    """
    Normalise WGS variant IDs from chr1:924533:A:G format to 1:924533:G
    - Strips 'chr' prefix
    - Drops REF allele, keeps CHROM:POS:ALT
    """
    parts = str(v).replace("chr", "").split(":")
    if len(parts) == 4:
        return f"{parts[0]}:{parts[1]}:{parts[3]}"
    return v

def normalise_rna(v):
    """
    Normalise RNA-seq variant IDs from 1:16257:C:23.8866 format to 1:16257:C
    - Drops QUAL score, keeps CHROM:POS:ALT
    """
    parts = str(v).split(":")
    if len(parts) >= 3:
        return f"{parts[0]}:{parts[1]}:{parts[2]}"
    return v

def save_outputs(overlap_df, tsv_path, xlsx_path):
    """Save overlap matrix to TSV and Excel with white->yellow->red colour scale."""
    overlap_df.to_csv(tsv_path, sep="\t")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Overlap Matrix"

    ws.cell(row=1, column=1, value="WGS \\ RNA")
    for j, col in enumerate(overlap_df.columns, start=2):
        ws.cell(row=1, column=j, value=col)
    for i, (idx, row) in enumerate(overlap_df.iterrows(), start=2):
        ws.cell(row=i, column=1, value=idx)
        for j, val in enumerate(row, start=2):
            ws.cell(row=i, column=j, value=int(val))

    data_range = f"B2:{get_column_letter(len(overlap_df.columns)+1)}{len(overlap_df)+1}"
    rule = ColorScaleRule(
        start_type="min", start_color="FFFFFF",
        mid_type="percentile", mid_value=50, mid_color="FFF59D",
        end_type="max", end_color="E53935"
    )
    ws.conditional_formatting.add(data_range, rule)

    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
    for row in ws.iter_rows(min_row=2, max_col=1):
        for cell in row:
            cell.font = Font(bold=True)

    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 2, 15)

    wb.save(xlsx_path)
    print(f"Saved {xlsx_path} and {tsv_path}")


# =============================================================================
# LOAD & NORMALISE DATA (shared across all analyses)
# =============================================================================

# =============================================================================
# FILE PATHS — edit these to point to your input files and output directory
# =============================================================================

WGS_FILE = "'/Users/giannabusch/RajLab Dropbox/Gianna Busch/Shared_GiannaB/Paper_Submission/WorkingPaper/Paper_DataandCode/ExtractedData/WholeGenomeSequencing/WGS_ExonicSNPs/wgs_comparison_table.txt'"
RNA_FILE  = "'/Users/giannabusch/RajLab Dropbox/Gianna Busch/Shared_GiannaB/Paper_Submission/WorkingPaper/Paper_DataandCode/ExtractedData/WholeGenomeSequencing/RNASeq_forCloneMatching/varianttxtfiles/rnaseq_variants_comparison_table.txt'"
OUTPUT_DIR = "'/Users/giannabusch/RajLab Dropbox/Gianna Busch/Shared_GiannaB/Paper_Submission/WorkingPaper/Paper_DataandCode/Plots'"          # e.g. "/home/user/results" or "C:/Users/me/results"

# =============================================================================

import os
os.makedirs(OUTPUT_DIR, exist_ok=True)

wgs_raw = pd.read_csv(WGS_FILE, sep="\t", index_col=0)
rna_raw = pd.read_csv(RNA_FILE,  sep="\t", index_col=0)

# Drop the RNA-seq format-description row (not a real variant)
rna_raw = rna_raw.drop(index="CHROM:POS:ALT:QUAL", errors="ignore")

# Normalise variant IDs to CHROM:POS:ALT format
wgs_raw.index = wgs_raw.index.map(normalise_wgs)
rna_raw.index = rna_raw.index.map(normalise_rna)

# Drop duplicate variant IDs (keep first occurrence)
wgs_raw = wgs_raw[~wgs_raw.index.duplicated(keep='first')]
rna_raw = rna_raw[~rna_raw.index.duplicated(keep='first')]

print(f"WGS: {wgs_raw.shape[0]} variants x {wgs_raw.shape[1]} samples")
print(f"RNA: {rna_raw.shape[0]} variants x {rna_raw.shape[1]} samples")


# =============================================================================
# MATRIX 1: All variants (no filtering)
# - Uses all variants common to both datasets
# =============================================================================
print("\n--- Matrix 1: No filtering ---")

common = sorted(set(wgs_raw.index) & set(rna_raw.index))
print(f"Common variants: {len(common)}")

wgs_bin = (wgs_raw.loc[common] > 0).astype(np.int16)
rna_bin = (rna_raw.loc[common] > 0).astype(np.int16)

overlap = wgs_bin.T.values @ rna_bin.values
overlap_df_1 = pd.DataFrame(overlap, index=wgs_bin.columns, columns=rna_bin.columns)

save_outputs(overlap_df_1, os.path.join(OUTPUT_DIR, "overlap_matrix.tsv"), os.path.join(OUTPUT_DIR, "overlap_matrix.xlsx"))


# =============================================================================
# MATRIX 2: Remove variants present in ALL WGS samples
# - Removes shared cancer cell line background mutations present in every sample
# =============================================================================
print("\n--- Matrix 2: Remove variants present in all WGS samples ---")

wgs_filtered = wgs_raw[~(wgs_raw > 0).all(axis=1)]
print(f"WGS variants after removing universal variants: {len(wgs_filtered)}")

common_2 = sorted(set(wgs_filtered.index) & set(rna_raw.index))
print(f"Common variants: {len(common_2)}")

wgs_bin = (wgs_filtered.loc[common_2] > 0).astype(np.int16)
rna_bin = (rna_raw.loc[common_2] > 0).astype(np.int16)

overlap = wgs_bin.T.values @ rna_bin.values
overlap_df_2 = pd.DataFrame(overlap, index=wgs_bin.columns, columns=rna_bin.columns)

save_outputs(overlap_df_2, os.path.join(OUTPUT_DIR, "overlap_matrix_filtered.tsv"), os.path.join(OUTPUT_DIR, "overlap_matrix_filtered.xlsx"))


# =============================================================================
# MATRIX 3: Variants present in exactly 1 WGS sample (unique/private variants)
# - Most discriminating but fewest variants; low RNA-seq capture rate
# =============================================================================
print("\n--- Matrix 3: Variants present in exactly 1 WGS sample ---")

wgs_unique = wgs_raw[(wgs_raw > 0).sum(axis=1) == 1]
print(f"WGS variants unique to a single sample: {len(wgs_unique)}")

common_3 = sorted(set(wgs_unique.index) & set(rna_raw.index))
print(f"Common variants: {len(common_3)}")

wgs_bin = (wgs_unique.loc[common_3] > 0).astype(np.int16)
rna_bin = (rna_raw.loc[common_3] > 0).astype(np.int16)

overlap = wgs_bin.T.values @ rna_bin.values
overlap_df_3 = pd.DataFrame(overlap, index=wgs_bin.columns, columns=rna_bin.columns)

save_outputs(overlap_df_3, os.path.join(OUTPUT_DIR, "overlap_matrix_unique.tsv"), os.path.join(OUTPUT_DIR, "overlap_matrix_unique.xlsx"))


# =============================================================================
# MATRIX 4: Variants present in 1-5 WGS samples (rare variants)
# - Middle ground: rare enough to be informative, more common than unique variants
# =============================================================================
print("\n--- Matrix 4: Variants present in <=5 WGS samples ---")

wgs_max5 = wgs_raw[(wgs_raw > 0).sum(axis=1) <= 5]
print(f"WGS variants present in <=5 samples: {len(wgs_max5)}")

common_4 = sorted(set(wgs_max5.index) & set(rna_raw.index))
print(f"Common variants: {len(common_4)}")

wgs_bin = (wgs_max5.loc[common_4] > 0).astype(np.int16)
rna_bin = (rna_raw.loc[common_4] > 0).astype(np.int16)

overlap = wgs_bin.T.values @ rna_bin.values
overlap_df_4 = pd.DataFrame(overlap, index=wgs_bin.columns, columns=rna_bin.columns)

save_outputs(overlap_df_4, os.path.join(OUTPUT_DIR, "overlap_matrix_max5.tsv"), os.path.join(OUTPUT_DIR, "overlap_matrix_max5.xlsx"))


# =============================================================================
# MATRIX 5: Exclusive variants per WGS sample
# - For each WGS sample, counts RNA-seq overlap using only variants that are
#   present in that WGS sample AND absent in ALL other WGS samples
# - Most stringent approach; very sparse due to low RNA-seq capture of rare variants
# =============================================================================
print("\n--- Matrix 5: Exclusive variants per WGS sample ---")

wgs_bin_full = (wgs_raw > 0).astype(int)
rna_bin_full = (rna_raw > 0).astype(int)

common_5 = sorted(set(wgs_bin_full.index) & set(rna_bin_full.index))
print(f"Common variants: {len(common_5)}")

wgs_common = wgs_bin_full.loc[common_5]
rna_common = rna_bin_full.loc[common_5]

results = {}
for sample in wgs_common.columns:
    present_in_sample = wgs_common[sample] == 1
    absent_in_others = (wgs_common.drop(columns=sample) == 0).all(axis=1)
    exclusive_variants = wgs_common.index[present_in_sample & absent_in_others]

    if len(exclusive_variants) == 0:
        results[sample] = pd.Series(0, index=rna_common.columns)
    else:
        results[sample] = rna_common.loc[exclusive_variants].sum(axis=0)

    print(f"  {sample}: {len(exclusive_variants)} exclusive variants")

overlap_df_5 = pd.DataFrame(results).T
overlap_df_5.index.name = "WGS \\ RNA"

save_outputs(overlap_df_5, os.path.join(OUTPUT_DIR, "overlap_matrix_exclusive.tsv"), os.path.join(OUTPUT_DIR, "overlap_matrix_exclusive.xlsx"))
